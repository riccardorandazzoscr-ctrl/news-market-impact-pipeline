#!/usr/bin/env python3
"""
export_to_excel.py

Crea una COPIA Excel leggibile del dataset di mercato (market_data.db).
NON tocca il DB: legge in sola lettura e scrive un nuovo file .xlsx datato in
market_data/exports/. Usalo per "dare un'occhiata" ai dati in Excel senza
rischiare di corrompere la fonte (che viene rigenerata ogni giorno da launchd).

Uso:
    venv/bin/python export_to_excel.py

Per un singolo asset / un intervallo, modifica le query in fondo o filtra in Excel.
"""

from datetime import date
from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from bootstrap_market_data import DB_PATH

EXPORT_DIR = Path(DB_PATH).parent / "exports"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    # larghezza colonne (cap a 40)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 40)


def main():
    conn = sqlite3.connect(DB_PATH)
    assets = pd.read_sql_query("SELECT * FROM assets ORDER BY asset_class, ticker", conn)
    prices = pd.read_sql_query("SELECT * FROM prices ORDER BY ticker, date", conn)

    # Riepilogo per ticker: copertura e ultimo prezzo
    summary = pd.read_sql_query(
        """
        SELECT p.ticker,
               a.name,
               a.asset_class,
               a.region,
               COUNT(*)                               AS n_righe,
               MIN(p.date)                            AS dal,
               MAX(p.date)                            AS al,
               ROUND((SELECT close FROM prices x
                      WHERE x.ticker = p.ticker
                      ORDER BY x.date DESC LIMIT 1), 2) AS ultimo_close
        FROM prices p
        LEFT JOIN assets a ON a.ticker = p.ticker
        GROUP BY p.ticker
        ORDER BY a.asset_class, p.ticker
        """,
        conn,
    )
    conn.close()

    guida = pd.DataFrame({
        "Foglio": ["Riepilogo", "Assets", "Prezzi"],
        "Cosa contiene": [
            "Una riga per asset: copertura date, numero di righe, ultimo prezzo. Parti da qui.",
            "Anagrafica dei 17 strumenti: nome, classe, area, valuta, fonte.",
            "Dati grezzi: una riga per (ticker, giorno) con OHLCV + adj_close. "
            "Usa il filtro in alto per isolare un ticker o un periodo.",
        ],
    })

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    out = EXPORT_DIR / f"market_data_snapshot_{date.today().isoformat()}.xlsx"
    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        guida.to_excel(xl, sheet_name="Guida", index=False)
        summary.to_excel(xl, sheet_name="Riepilogo", index=False)
        assets.to_excel(xl, sheet_name="Assets", index=False)
        prices.to_excel(xl, sheet_name="Prezzi", index=False)
        for name in ("Guida", "Riepilogo", "Assets", "Prezzi"):
            _style_header(xl.book[name])

    print(f"[export] scritto {out}  ({len(prices):,} righe prezzi, {len(assets)} asset)")


if __name__ == "__main__":
    main()
